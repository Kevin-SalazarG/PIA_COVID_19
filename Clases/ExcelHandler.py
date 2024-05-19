import openpyxl
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from Utils.TextColor import TextColor

class ExcelHandler:
    def __init__(self, filename: str):
        self.filename = filename
        self.wb = Workbook()
        self.ws = self.wb.active

    def create_sheet(self, title: str):
        self.ws.title = title

    def add_headers(self, headers: list):
        self.ws.append(headers)

    def add_row(self, row: list):
        self.ws.append(row)

    def adjust_column_widths(self):
        for col in self.ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = max_length + 2
            self.ws.column_dimensions[column].width = adjusted_width

    def create_bar_chart(self, title: str, x_title: str, y_title: str, data_range, category_range, position: str):
        chart = BarChart()
        chart.title = title
        chart.x_axis.title = x_title
        chart.y_axis.title = y_title
        chart.add_data(data_range, titles_from_data=True)
        chart.set_categories(category_range)
        self.ws.add_chart(chart, position)

    def save(self):
        self.wb.save(self.filename)
        print(TextColor.GREEN, "Archivo guardado como", TextColor.WHITE, self.filename)
